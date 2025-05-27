import streamlit as st
import pandas as pd
import pdfplumber
import docx
from dotenv import load_dotenv
import os
import json
import time
import nltk
import spacy
from spacy.util import get_installed_models
from cachetools import TTLCache
from enum import Enum
from dataclasses import dataclass
from google.api_core.exceptions import ResourceExhausted
import google.generativeai as genai
from langchain_google_genai import ChatGoogleGenerativeAI
from langchain_core.prompts import ChatPromptTemplate
from langchain_core.output_parsers import StrOutputParser
import logging
import re

# Configure logging
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)
logging.getLogger("pdfminer").setLevel(logging.ERROR)

# Initialize caches
resume_cache = TTLCache(maxsize=50, ttl=3600)
cover_letter_cache = TTLCache(maxsize=50, ttl=3600)
skills_gap_cache = TTLCache(maxsize=50, ttl=3600)
interview_cache = TTLCache(maxsize=50, ttl=3600)

# Load environment variables
load_dotenv()

# Agent States
class AgentState(Enum):
    IDLE = "idle"
    PROCESSING_RESUME = "processing_resume"
    GENERATING_COVER_LETTER = "generating_cover_letter"
    ANALYZING_SKILLS_GAP = "analyzing_skills_gap"
    GENERATING_INTERVIEW = "generating_interview"

# User Profile Data Structure
@dataclass
class UserProfile:
    name: str
    email: str
    phone: str
    address: str
    education: list
    skills: list
    experience: list
    publications: list
    job_role: str = None

# Mock course database
COURSE_DATABASE = {
    "Python": [
        {"title": "Python for Data Science and Machine Learning Bootcamp", "platform": "Udemy", "url": "https://www.udemy.com/course/python-for-data-science-and-machine-learning-bootcamp/", "duration": "25 hours"},
        {"title": "Python Programming Masterclass", "platform": "Coursera", "url": "https://www.coursera.org/learn/python-programming", "duration": "4 weeks"}
    ],
    "SQL": [
        {"title": "SQL for Data Analysis", "platform": "Udemy", "url": "https://www.udemy.com/course/the-complete-sql-bootcamp/", "duration": "9 hours"},
        {"title": "SQL and Relational Databases 101", "platform": "Coursera", "url": "https://www.coursera.org/learn/sql-and-relational-databases", "duration": "3 weeks"}
    ],
    "Java": [
        {"title": "Java Programming and Software Engineering Fundamentals", "platform": "Coursera", "url": "https://www.coursera.org/specializations/java-programming", "duration": "5 months"},
        {"title": "Java Masterclass - Beginner to Expert", "platform": "Udemy", "url": "https://www.udemy.com/course/java-the-complete-java-developer-course/", "duration": "68 hours"}
    ],
    "Communication": [
        {"title": "Effective Communication: Writing, Design, and Presentation", "platform": "Coursera", "url": "https://www.coursera.org/specializations/effective-communication", "duration": "2 months"}
    ],
    "Data Analysis": [
        {"title": "Data Analysis with Python", "platform": "Coursera", "url": "https://www.coursera.org/learn/data-analysis-with-python", "duration": "4 weeks"},
        {"title": "Excel to MySQL: Analytic Techniques for Business", "platform": "Coursera", "url": "https://www.coursera.org/specializations/excel-mysql", "duration": "5 months"}
    ],
    "Machine Learning": [
        {"title": "Machine Learning by Stanford Online", "platform": "Coursera", "url": "https://www.coursera.org/learn/machine-learning", "duration": "11 weeks"},
        {"title": "Deep Learning Specialization", "platform": "Coursera", "url": "https://www.coursera.org/specializations/deep-learning", "duration": "4 months"}
    ],
    "ETL": [
        {"title": "Data Engineering Fundamentals", "platform": "Coursera", "url": "https://www.coursera.org/learn/data-engineering-fundamentals", "duration": "4 weeks"},
        {"title": "ETL and Data Pipelines with Python", "platform": "Udemy", "url": "https://www.udemy.com/course/etl-and-data-pipelines-with-python/", "duration": "12 hours"}
    ],
    "Cloud Computing": [
        {"title": "AWS Cloud Practitioner Essentials", "platform": "Coursera", "url": "https://www.coursera.org/learn/aws-cloud-practitioner-essentials", "duration": "3 weeks"},
        {"title": "Google Cloud Platform Fundamentals", "platform": "Coursera", "url": "https://www.coursera.org/learn/gcp-fundamentals", "duration": "2 weeks"}
    ]
}

# Job Seeker Agent
class JobSeekerAgent:
    def __init__(self):
        self.state = AgentState.IDLE
        self.llm = self._init_llm()
        self.nlp = self._init_nlp()
        self.user_profile = None
        self.sri_lankan_context = self._load_sri_lankan_context()

    def _init_llm(self):
        api_key = os.getenv("GOOGLE_API_KEY")
        if not api_key:
            logger.error("GOOGLE_API_KEY not found in .env file.")
            st.error("GOOGLE_API_KEY not found. Add it to your .env file.")
            return None
        try:
            genai.configure(api_key=api_key)
            llm = ChatGoogleGenerativeAI(
                model="gemini-1.5-flash",
                google_api_key=api_key,
                temperature=0.7,
                max_tokens=2048
            )
            logger.info("LLM initialized successfully.")
            return llm
        except Exception as e:
            logger.error(f"Failed to initialize LLM: {e}")
            st.error(f"Failed to initialize LLM: {e}")
            return None

    def _init_nlp(self):
        if "en_core_web_sm" not in get_installed_models():
            logger.error("spaCy model 'en_core_web_sm' not found.")
            st.error("spaCy model 'en_core_web_sm' not found. Run: `python -m spacy download en_core_web_sm`")
            return None
        try:
            nlp = spacy.load("en_core_web_sm")
            logger.info("spaCy model loaded successfully.")
            return nlp
        except Exception as e:
            logger.error(f"Error loading spaCy model: {e}")
            st.error(f"Error loading spaCy model: {e}")
            return None

    def _load_sri_lankan_context(self):
        return {
            "major_cities": ["Colombo", "Kandy", "Galle", "Jaffna"],
            "top_industries": ["IT & Software", "Banking & Finance", "Tourism", "Healthcare"],
            "popular_companies": ["Dialog", "Virtusa", "WSO2", "IFS"],
            "job_portals": ["topjobs.lk", "jobs.lk", "ikman.lk", "linkedin.com"]
        }

    def set_state(self, new_state: AgentState):
        logger.info(f"Agent state changing from {self.state} to {new_state}")
        self.state = new_state

    def process_resume(self, resume_text: str, job_role: str, retries=3, delay=44):
        self.set_state(AgentState.PROCESSING_RESUME)
        if not self.llm:
            logger.error("LLM not initialized during resume processing.")
            st.error("LLM not initialized. Check GOOGLE_API_KEY.")
            return {}
        
        prompt = ChatPromptTemplate.from_template("""
        You are an expert in resume analysis for the Sri Lankan job market. Perform the following tasks:
        1. Extract details from the resume:
           - Name
           - Email
           - Phone (Sri Lankan format, e.g., +94 or 0 followed by 9 digits)
           - Address
           - Education (degrees, institutions, years)
           - Skills (technical and soft skills)
           - Experience (job titles, companies, durations, descriptions)
           - Publications (if any)
        2. Rate the resume for the job role '{job_role}' based on clarity, relevance, and completeness (0-100%) with a brief explanation.
        3. Suggest 3-5 job roles suitable for the candidate based on their skills and experience.
        Output ONLY a valid JSON object with keys: 'details', 'rating', 'job_roles'. Do not include any text outside the JSON.
        Example:
        {{
            "details": {{
                "name": "John Doe",
                "email": "john@example.com",
                "phone": "+94123456789",
                "address": "123 Main St, Colombo",
                "education": [{{"degree": "BSc", "institution": "University of Colombo", "year": "2018"}}],
                "skills": ["Python", "SQL", "Communication"],
                "experience": [{{"title": "Software Engineer", "company": "WSO2", "duration": "2018-2020", "description": "Developed APIs"}}],
                "publications": []
            }},
            "rating": {{
                "score": 85,
                "explanation": "Clear and relevant experience."
            }},
            "job_roles": ["Software Engineer", "Data Analyst", "DevOps Engineer"]
        }}
        Resume Text: {resume_text}
        """)
        chain = prompt | self.llm | StrOutputParser()
        cache_key = f"resume_{hash(resume_text)}_{job_role}"
        
        if cache_key in resume_cache:
            logger.info("Resume found in cache.")
            return resume_cache[cache_key]

        result = {}
        for attempt in range(retries):
            try:
                logger.info(f"Attempting resume processing (attempt {attempt + 1}/{retries})")
                raw_result = chain.invoke({"resume_text": resume_text, "job_role": job_role})
                logger.debug(f"Raw API response: {raw_result}")
                
                cleaned_result = re.sub(r'```json\s*|\s*```', '', raw_result).strip()
                data = json.loads(cleaned_result)
                
                if not all(key in data for key in ['details', 'rating', 'job_roles']):
                    raise ValueError("Response missing required keys: 'details', 'rating', 'job_roles'")
                
                resume_cache[cache_key] = data
                self.user_profile = UserProfile(
                    name=data['details'].get('name', ''),
                    email=data['details'].get('email', ''),
                    phone=data['details'].get('phone', ''),
                    address=data['details'].get('address', ''),
                    education=data['details'].get('education', []),
                    skills=data['details'].get('skills', []),
                    experience=data['details'].get('experience', []),
                    publications=data['details'].get('publications', []),
                    job_role=job_role
                )
                st.session_state.user_profile = self.user_profile
                logger.info(f"Resume processed successfully, user profile set for job role: {job_role}")
                return data
            except ResourceExhausted:
                logger.warning(f"Quota exceeded (attempt {attempt + 1}/{retries}). Retrying in {delay}s...")
                st.warning(f"Quota exceeded (attempt {attempt + 1}/{retries}). Retrying in {delay}s...")
                time.sleep(delay)
                delay *= 2
            except json.JSONDecodeError as e:
                logger.error(f"JSON parsing error: {e}, response: {cleaned_result}")
                st.error(f"JSON parsing error: {e}")
                return {}
            except Exception as e:
                logger.error(f"Error processing resume: {e}")
                st.error(f"Error processing resume: {e}")
                return {}
        logger.error("Failed to process resume due to quota limits.")
        st.error("Failed to process resume due to quota limits. Check https://ai.google.dev/gemini-api/docs/rate-limits.")
        return {}

    def generate_cover_letter(self, profile: UserProfile, retries=3, delay=44):
        self.set_state(AgentState.GENERATING_COVER_LETTER)
        if not self.llm:
            logger.error("LLM not initialized for cover letter generation.")
            st.error("LLM not initialized. Check GOOGLE_API_KEY.")
            return ""
        
        if not profile or not profile.job_role:
            logger.error("User profile or job role missing.")
            st.error("User profile or job role missing. Analyze a resume first.")
            return ""

        logger.info(f"Generating cover letter for job role: {profile.job_role}")
        
        cache_key = f"cover_letter_{hash(str(profile.skills) + profile.job_role)}"
        if cache_key in cover_letter_cache:
            logger.info("Cover letter found in cache.")
            return cover_letter_cache[cache_key]

        prompt = ChatPromptTemplate.from_template("""
        You are a career consultant crafting a professional cover letter (300-400 words) for a candidate applying for a {job_role} position in Sri Lanka. Use the candidate's CV details to tailor the letter. Highlight the candidate's most relevant skills, experiences, and education that align with typical requirements for the job role. Ensure a formal, enthusiastic tone and focus on how the candidate's background makes them a strong fit for the position.

        Candidate CV Details:
        - Name: {name}
        - Email: {email}
        - Skills: {skills}
        - Experience: {experience}
        - Education: {education}
        - Job Role: {job_role}

        Output the cover letter as plain text, formatted with proper salutation (e.g., "Dear Hiring Manager") and closing (e.g., "Sincerely, [Name]"). Do not include placeholders like [Your Name] or [Company Name].
        """)
        chain = prompt | self.llm | StrOutputParser()

        for attempt in range(retries):
            try:
                logger.info(f"Attempting cover letter generation (attempt {attempt + 1}/{retries})")
                result = chain.invoke({
                    "name": profile.name,
                    "email": profile.email,
                    "skills": ", ".join(profile.skills),
                    "experience": "; ".join([f"{exp.get('title', '')} at {exp.get('company', '')} ({exp.get('duration', '')}): {exp.get('description', '')}" for exp in profile.experience]),
                    "education": "; ".join([f"{edu.get('degree', '')} from {edu.get('institution', '')} ({edu.get('year', '')})" for edu in profile.education]),
                    "job_role": profile.job_role
                })
                logger.info("Cover letter generated successfully.")
                cover_letter_cache[cache_key] = result
                return result
            except ResourceExhausted:
                logger.warning(f"Quota exceeded (attempt {attempt + 1}/{retries}). Retrying in {delay}s...")
                st.warning(f"Quota exceeded (attempt {attempt + 1}/{retries}). Retrying in {delay}s...")
                time.sleep(delay)
                delay *= 2
            except Exception as e:
                logger.error(f"Error generating cover letter: {e}")
                st.error(f"Error generating cover letter: {e}")
                return ""
        logger.error("Failed to generate cover letter due to quota limits.")
        st.error("Failed to generate cover letter due to quota limits.")
        return ""

    def skills_gap_analysis(self, profile: UserProfile, retries=3, delay=44):
        self.set_state(AgentState.ANALYZING_SKILLS_GAP)
        if not self.llm:
            logger.error("LLM not initialized for skills gap analysis.")
            st.error("LLM not initialized.")
            return [], []

        cache_key = f"skills_gap_{hash(str(profile.skills) + profile.job_role)}"
        if cache_key in skills_gap_cache:
            logger.info("Skills gap analysis found in cache.")
            return skills_gap_cache[cache_key]

        prompt = ChatPromptTemplate.from_template("""
        You are a career advisor analyzing the skills gap for a {job_role} position in Sri Lanka. Given the candidate's skills and the job role, perform the following:
        1. Generate a list of recommended skills (technical and soft) typically required for the {job_role} in the Sri Lankan job market (strictly 8-12 skills for relevance).
        2. Compare these with the candidate's skills to identify:
           - Matched skills (skills the candidate has that are required).
           - Missing skills (required skills the candidate lacks).
        Output a JSON object with 'matched_skills' and 'missing_skills' as lists. Ensure missing_skills does not exceed 10.

        Candidate Skills: {candidate_skills}
        Job Role: {job_role}

        Example:
        {{
            "matched_skills": ["Python", "SQL"],
            "missing_skills": ["Machine Learning", "Cloud Computing"]
        }}
        """)
        chain = prompt | self.llm | StrOutputParser()

        for attempt in range(retries):
            try:
                logger.info(f"Attempting skills gap analysis (attempt {attempt + 1}/{retries})")
                result = chain.invoke({
                    "candidate_skills": ", ".join(profile.skills),
                    "job_role": profile.job_role
                })
                cleaned_result = re.sub(r'```json\s*|\s*```', '', result).strip()
                data = json.loads(cleaned_result)
                
                if not all(key in data for key in ['matched_skills', 'missing_skills']):
                    raise ValueError("Response missing required keys: 'matched_skills', 'missing_skills'")
                
                # Enforce missing skills limit
                data['missing_skills'] = data['missing_skills'][:10]
                
                skills_gap_cache[cache_key] = (data['matched_skills'], data['missing_skills'])
                logger.info("Skills gap analysis completed.")
                return data['matched_skills'], data['missing_skills']
            except ResourceExhausted:
                logger.warning(f"Quota exceeded (attempt {attempt + 1}/{retries}). Retrying in {delay}s...")
                st.warning(f"Quota exceeded (attempt {attempt + 1}/{retries}). Retrying in {delay}s...")
                time.sleep(delay)
                delay *= 2
            except json.JSONDecodeError as e:
                logger.error(f"JSON parsing error: {e}, response: {cleaned_result}")
                st.error(f"JSON parsing error: {e}")
                return [], []
            except Exception as e:
                logger.error(f"Error in skills gap analysis: {e}")
                st.error(f"Error in skills gap analysis: {e}")
                return [], []
        logger.error("Failed to perform skills gap analysis due to quota limits.")
        st.error("Failed to perform skills gap analysis due to quota limits.")
        return [], []

    def recommend_learning_path(self, missing_skills: list):
        roadmap = []
        for week, skill in enumerate(missing_skills[:4], 1):
            if skill in COURSE_DATABASE:
                course = COURSE_DATABASE[skill][0]
                roadmap.append({
                    "week": week,
                    "skill": skill,
                    "course": course["title"],
                    "platform": course["platform"],
                    "url": course["url"],
                    "duration": course["duration"]
                })
            else:
                roadmap.append({
                    "week": week,
                    "skill": skill,
                    "course": f"Learn {skill}",
                    "platform": "Various",
                    "url": f"https://www.coursera.org/search?query={skill.replace(' ', '%20')}",
                    "duration": "Varies"
                })
        logger.info(f"Learning path recommended for {len(roadmap)} skills.")
        return roadmap

    def generate_interview(self, profile: UserProfile, retries=3, delay=44):
        self.set_state(AgentState.GENERATING_INTERVIEW)
        if not self.llm:
            logger.error("LLM not initialized for interview generation.")
            st.error("LLM not initialized.")
            return []

        cache_key = f"interview_{hash(profile.job_role)}"
        if cache_key in interview_cache:
            logger.info("Interview questions found in cache.")
            return interview_cache[cache_key]

        prompt = ChatPromptTemplate.from_template("""
        You are a hiring manager for a {job_role} position in Sri Lanka. Generate 5 interview questions (3 technical, 2 behavioral) tailored to the job role and the candidate's CV details. The questions should assess the candidate's skills, experience, and education as they relate to typical requirements for the job role. Return a JSON array of questions with 'type' (technical or behavioral) and 'question' fields.

        Candidate CV Details:
        - Skills: {skills}
        - Experience: {experience}
        - Education: {education}
        - Job Role: {job_role}

        Example:
        [
            {{"type": "technical", "question": "Explain how you would optimize a SQL query for a large dataset."}},
            {{"type": "behavioral", "question": "Describe a time you led a team to meet a tight project deadline."}}
        ]
        """)
        chain = prompt | self.llm | StrOutputParser()

        for attempt in range(retries):
            try:
                logger.info(f"Attempting interview question generation (attempt {attempt + 1}/{retries})")
                result = chain.invoke({
                    "skills": ", ".join(profile.skills),
                    "experience": "; ".join([f"{exp.get('title', '')} at {exp.get('company', '')} ({exp.get('duration', '')}): {exp.get('description', '')}" for exp in profile.experience]),
                    "education": "; ".join([f"{edu.get('degree', '')} from {edu.get('institution', '')} ({edu.get('year', '')})" for edu in profile.education]),
                    "job_role": profile.job_role
                })
                result = result.strip()
                if result.startswith('```json'):
                    result = result[7:-3].strip()
                json_match = re.search(r'\[\s*\{.*?\}\s*\]', result, re.DOTALL)
                if not json_match:
                    raise ValueError(f"No valid JSON array found in response: {result}")
                questions = json.loads(json_match.group(0))
                for q in questions:
                    if not all(k in q for k in ['type', 'question']):
                        raise ValueError(f"Invalid question format: {q}")
                interview_cache[cache_key] = questions
                logger.info("Interview questions generated successfully.")
                return questions
            except ResourceExhausted:
                logger.warning(f"Quota exceeded (attempt {attempt + 1}/{retries}). Retrying in {delay}s...")
                st.warning(f"Quota exceeded (attempt {attempt + 1}/{retries}). Retrying in {delay}s...")
                time.sleep(delay)
                delay *= 2
            except Exception as e:
                logger.error(f"Error generating interview questions: {e}")
                st.error(f"Error generating interview questions: {e}")
                return []
        logger.error("Failed to generate interview questions due to quota limits.")
        return []

    def analyze_interview_response(self, question: str, response: str, retries=3, delay=44):
        if not self.llm:
            logger.error("LLM not initialized for interview response analysis.")
            st.error("LLM not initialized.")
            return {"score": 0, "feedback": "Analysis failed."}

        prompt = ChatPromptTemplate.from_template("""
        You are a hiring manager evaluating an interview response for a {job_role} position. Analyze the response to:

        Question: {question}
        Response: {response}

        Return a JSON object with:
        - 'score' (0-100): Based on relevance, clarity, and keyword usage.
        - 'feedback': Suggestions for improvement.

        Example:
        {{
            "score": 85,
            "feedback": "Good example, include more technical details."
        }}
        """)
        chain = prompt | self.llm | StrOutputParser()

        for attempt in range(retries):
            try:
                logger.info(f"Attempting interview response analysis (attempt {attempt + 1}/{retries})")
                result = chain.invoke({"job_role": self.user_profile.job_role, "question": question, "response": response})
                json_match = re.search(r'\{.*?\}(?=\s*$|\n|$)', result, re.DOTALL)
                if not json_match:
                    raise ValueError("No valid JSON found")
                data = json.loads(json_match.group(0))
                logger.info("Interview response analyzed successfully.")
                return data
            except ResourceExhausted:
                logger.warning(f"Quota exceeded (attempt {attempt + 1}/{retries}). Retrying in {delay}s...")
                st.warning(f"Quota exceeded (attempt {attempt + 1}/{retries}). Retrying in {delay}s...")
                time.sleep(delay)
                delay *= 2
            except Exception as e:
                logger.error(f"Error analyzing response: {e}")
                st.error(f"Error analyzing response: {e}")
                return {"score": 0, "feedback": "Analysis failed."}
        logger.error("Failed to analyze response due to quota limits.")
        return {"score": 0, "feedback": "Analysis failed due to quota limits."}

def init_nltk():
    resources = ['punkt', 'averaged_perceptron_tagger', 'maxent_ne_chunker', 'words']
    for resource in resources:
        try:
            nltk.download(resource, quiet=True)
            logger.info(f"NLTK resource {resource} downloaded successfully.")
        except Exception as e:
            logger.warning(f"Failed to download NLTK resource {resource}: {e}")

def main():
    # Check for PyTorch
    try:
        __import__("torch")
        st.error("PyTorch detected. Please uninstall it to avoid conflicts. Run: pip uninstall torch torchvision torchaudio -y")
        logger.error("PyTorch detected in environment.")
        return
    except ImportError:
        logger.info("No PyTorch detected, proceeding.")

    st.set_page_config(page_title="Sri Lankan Job Seeker AI", layout="wide", page_icon="🇱🇰")
    st.markdown("""
    <style>
    .main-header {
        background: linear-gradient(45deg, #1e3c72 0%, #3a7bd5 100%);
        padding: 2rem;
        border-radius: 8px;
        color: white;
        text-align: center;
        margin-bottom: 1.5rem;
    }
    </style>
    """, unsafe_allow_html=True)

    init_nltk()

    if not os.getenv("GOOGLE_API_KEY"):
        st.error("GOOGLE_API_KEY not found in .env file. Please add it and restart the app.")
        logger.error("GOOGLE_API_KEY not found in .env file.")
        return

    if 'agent' not in st.session_state:
        try:
            st.session_state.agent = JobSeekerAgent()
            logger.info("Agent initialized in session state.")
        except Exception as e:
            logger.error(f"Failed to initialize agent: {e}")
            st.error(f"Failed to initialize agent: {e}")
            return
    agent = st.session_state.agent

    if 'user_profile' in st.session_state and st.session_state.user_profile:
        agent.user_profile = st.session_state.user_profile
        logger.info(f"User profile restored from session state for job role: {agent.user_profile.job_role}")

    st.markdown("""
    <div class="main-header">
        <h1>🇱🇰 Sri Lankan Job Seeker AI Agent</h1>
        <p>Your intelligent career companion for the Sri Lankan job market</p>
    </div>
    """, unsafe_allow_html=True)

    with st.sidebar:
        st.header("🎯 Agent Status")
        st.write(f"**Current State**: {agent.state.value.title()}")
        st.header("🔧 Configuration")
        st.markdown("""
        1. Create a `.env` file with:
            ```bash
            GOOGLE_API_KEY=your_gemini_api_key
            ```
        2. Install libraries in a virtual environment:
            ```bash
            python -m venv venv
            .\venv\Scripts\activate
            pip install streamlit google-generativeai langchain==0.2.0 langchain-google-genai==1.0.5 pdfplumber python-docx spacy python-dotenv cachetools nltk pandas
            python -m spacy download en_core_web_sm
            ```
        3. Run with file watcher disabled:
            ```bash
            streamlit run cv_analyzer.py --server.fileWatcherType=none
            ```
        4. Check quotas: https://ai.google.dev/gemini-api/docs/rate-limits
        5. Verify spaCy: `python -m spacy validate`
        6. Ensure PyTorch is uninstalled:
            ```bash
            pip uninstall torch torchvision torchaudio -y
            pip list | findstr torch
            ```
            If output is empty, PyTorch is removed. If persistent, use:
            ```bash
            pip install pip-autoremove
            pip-autoremove torch -y
            ```
        7. Check Streamlit version (1.39.0+ required):
            ```bash
            pip show streamlit
            ```
            Upgrade if needed:
            ```bash
            pip install --upgrade streamlit
            ```
        8. Verify Python environment:
            ```bash
            where python
            ```
            Use: `C:\\Users\\sande\\AppData\\Local\\Programs\\Python\\Python313\\python.exe`
        """)

    tab1, tab2, tab3, tab4 = st.tabs(["📄 Resume Analysis", "✍️ Cover Letter", "📊 Skills Gap", "🎤 Mock Interview"])

    with tab1:
        st.header("📄 Resume Analysis")
        uploaded_file = st.file_uploader("Upload your CV (PDF, DOC, DOCX)", type=["pdf", "doc", "docx"])
        job_role = st.text_input("Desired Job Role", placeholder="e.g., Data Engineer")

        if st.button("Analyze CV", type="primary"):
            if not uploaded_file or not job_role:
                st.error("Please upload a CV and specify a job role.")
                logger.error("CV or job role missing.")
            else:
                with st.spinner("Analyzing resume..."):
                    file_extension = uploaded_file.name.split('.')[-1].lower()
                    resume_text = None
                    if file_extension == "pdf":
                        with pdfplumber.open(uploaded_file) as pdf:
                            resume_text = "".join(page.extract_text() or "" for page in pdf.pages)
                    elif file_extension in ["doc", "docx"]:
                        doc = docx.Document(uploaded_file)
                        resume_text = "\n".join(para.text for para in doc.paragraphs)

                    if resume_text:
                        result = agent.process_resume(resume_text, job_role)
                        if result:
                            st.subheader("Extracted Details")
                            if result.get('details'):
                                st.json(result['details'])
                            else:
                                st.warning("No details extracted.")
                                logger.warning("No details extracted from resume.")

                            st.subheader("CV Rating")
                            rating = result.get('rating', {})
                            st.write(f"**Score**: {rating.get('score', 0)}%")
                            st.write(f"**Explanation**: {rating.get('explanation', 'N/A')}")

                            st.subheader("Suggested Job Roles")
                            job_roles = result.get('job_roles', [])
                            if job_roles:
                                for role in job_roles:
                                    st.write(f"- {role}")
                            else:
                                st.warning("No job roles suggested.")
                                logger.warning("No job roles suggested.")
                        else:
                            logger.error("Resume processing returned no result.")
                    else:
                        st.error("Failed to extract text from CV.")
                        logger.error("Failed to extract text from CV.")

    with tab2:
        st.header("✍️ Cover Letter Generator")
        if agent.user_profile:
            st.info(f"Generating cover letter for {agent.user_profile.job_role}")
            logger.info(f"User profile available for cover letter, job role: {agent.user_profile.job_role}")
            if st.button("Generate Cover Letter", type="primary"):
                with st.spinner("Generating cover letter..."):
                    cover_letter = agent.generate_cover_letter(agent.user_profile)
                    if cover_letter:
                        edited_letter = st.text_area("Edit Cover Letter", value=cover_letter, height=500)
                        st.download_button(
                            label="Download Cover Letter",
                            data=edited_letter,
                            file_name="cover_letter.txt",
                            mime="text/plain"
                        )
                        logger.info("Cover letter displayed and available for download.")
                    else:
                        st.error("Failed to generate cover letter. Check logs for details.")
                        logger.error("Failed to generate cover letter.")
        else:
            st.info("Please analyze a resume first.")
            logger.info("No user profile available for cover letter.")

    with tab3:
        st.header("📊 Skills Gap Analysis")
        if agent.user_profile:
            st.info(f"Analyzing skills gap for {agent.user_profile.job_role}")
            if st.button("Analyze Skills Gap", type="primary"):
                with st.spinner("Analyzing skills gap..."):
                    matched_skills, missing_skills = agent.skills_gap_analysis(agent.user_profile)
                    if not matched_skills and not missing_skills:
                        st.subheader("No skills detected for the job role.")
                        logger.info("No skills detected for job role.")
                    else:
                        st.subheader("Skills Analysis")
                        st.write(f"**Matched Skills**: {', '.join(matched_skills) if matched_skills else 'None'}")
                        st.write(f"**Missing Skills**: {', '.join(missing_skills) if missing_skills else 'None'}")

                        st.subheader("Skills Gap Chart")
                        chart_data = pd.DataFrame({
                            "Category": ["Matched Skills", "Missing Skills"],
                            "Count": [len(matched_skills), len(missing_skills)]
                        }).set_index("Category")
                        st.bar_chart(chart_data, height=400)
                        logger.info("Bar chart rendered.")

                        st.subheader("Recommended Courses to Bridge Skills Gap")
                        roadmap = agent.recommend_learning_path(missing_skills)
                        if roadmap:
                            for item in roadmap:
                                st.write(f"**Week {item['week']}**: {item['skill']} - {item['course']} ({item['platform']}, {item['duration']})")
                                st.markdown(f"[Enroll]({item['url']})")
                        else:
                            st.warning("No courses found for missing skills.")
                            logger.warning("No courses found for missing skills.")
        else:
            st.info("Please analyze a resume first.")
            logger.info("No user profile available for skills gap analysis.")

    with tab4:
        st.header("🎤 AI Mock Interview")
        if agent.user_profile:
            st.info(f"Generating interview questions for {agent.user_profile.job_role}")
            logger.info(f"User profile available for mock interview, job role: {agent.user_profile.job_role}")
            
            # Initialize session state for interview
            if 'interview_questions' not in st.session_state:
                st.session_state.interview_questions = []
            if 'interview_responses' not in st.session_state:
                st.session_state.interview_responses = {}
            
            if st.button("Start Interview", type="primary"):
                with st.spinner("Generating interview questions..."):
                    questions = agent.generate_interview(agent.user_profile)
                    if questions:
                        st.session_state.interview_questions = questions
                        st.session_state.interview_responses = {str(i): {"response": "", "feedback": None} for i in range(len(questions))}
                        logger.info("Interview questions generated and stored in session state.")
                    else:
                        st.error("Failed to generate interview questions.")
                        logger.error("Failed to generate interview questions.")

            if st.session_state.interview_questions:
                st.subheader("Interview Session")
                for idx, q in enumerate(st.session_state.interview_questions):
                    st.write(f"**Question {idx + 1} ({q['type'].title()})**: {q['question']}")
                    
                    # Get or set response
                    response_key = str(idx)
                    response = st.text_area(
                        f"Your response to question {idx + 1}",
                        value=st.session_state.interview_responses[response_key]["response"],
                        key=f"q{idx}",
                        height=150
                    )
                    
                    # Update response in session state
                    st.session_state.interview_responses[response_key]["response"] = response

                    # Submit button
                    if st.button(f"Submit Response {idx + 1}", key=f"submit_q{idx}"):
                        if response.strip():
                            with st.spinner("Analyzing response..."):
                                feedback = agent.analyze_interview_response(q['question'], response)
                                st.session_state.interview_responses[response_key]["feedback"] = feedback
                                logger.info(f"Response to question {idx + 1} analyzed.")

                    # Display feedback if available
                    feedback = st.session_state.interview_responses[response_key]["feedback"]
                    if feedback:
                        status = "OK" if feedback['score'] >= 70 else "Needs Improvement"
                        st.write(f"**Status**: {status}")
                        st.write(f"**Score**: {feedback['score']}/100")
                        st.write(f"**Feedback**: {feedback['feedback']}")

                # Display summary if all questions answered
                if all(r["feedback"] for r in st.session_state.interview_responses.values()):
                    st.subheader("Interview Summary")
                    total_score = sum(r["feedback"]['score'] for r in st.session_state.interview_responses.values()) / len(st.session_state.interview_questions)
                    st.write(f"**Overall Score**: {round(total_score, 2)}/100")
                    st.write("**Improvement Suggestions**:")
                    for idx, q in enumerate(st.session_state.interview_questions):
                        feedback = st.session_state.interview_responses[str(idx)]["feedback"]
                        st.write(f"- **Q{idx + 1}**: {q['question']}")
                        st.write(f"  **Feedback**: {feedback['feedback']}")
                        status = "OK" if feedback['score'] >= 70 else "Needs Improvement"
                        st.write(f"  **Status**: {status}")
        else:
            st.info("Please analyze a resume first.")
            logger.info("No user profile available for mock interview.")

if __name__ == "__main__":
    main()